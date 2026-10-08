import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Title
title = doc.add_heading('MedPal Project Architecture & Tech Stack Report', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

# 1. Current Tech Stack
doc.add_heading('1. Current Technology Stack', level=1)
p1 = doc.add_paragraph()
p1.add_run('The current MedPal application is built using the following technologies:\n').bold = True
doc.add_paragraph('Frontend Framework: React (initialized via Vite). This is a modern, fast JavaScript framework for building user interfaces. We are using TypeScript (the .tsx files) which adds static typing to JavaScript for better error checking.')
doc.add_paragraph('Styling: Vanilla CSS (Standard CSS). We are currently using pure CSS with CSS variables for our design system, rather than a utility framework like Tailwind.')
doc.add_paragraph('Routing: react-router-dom. Used for navigating between pages (Dashboard, Fitness, AI Companion, etc.) without reloading the browser.')
doc.add_paragraph('Backend: Node.js with Express. A lightweight server to handle API requests and communicate with the AI.')
doc.add_paragraph('AI Integration: @google/genai SDK to connect directly to the Gemini 2.5 Flash model.')
doc.add_paragraph('Local Database: localForage (IndexedDB wrapper) for offline-first storage of health records on the browser.')

# 2. Alternative Approaches & Tailwind CSS
doc.add_heading('2. Comparison of Styling Alternatives (Vanilla CSS vs Tailwind CSS)', level=1)
doc.add_paragraph('Currently, we are using Vanilla CSS (like App.css, index.css). While this offers complete custom control, it can become hard to maintain as the project grows.')
doc.add_heading('Tailwind CSS (Utility-First CSS):', level=2)
doc.add_paragraph('Pros: Extremely fast to write once you learn it. You style elements directly in the HTML/JSX (e.g., <div className="bg-blue-500 text-white p-4 rounded-lg">). It enforces a consistent design system, makes dark mode trivial to implement, and eliminates the need to switch back and forth between .tsx and .css files.')
doc.add_paragraph('Cons: The HTML can look messy/cluttered with many class names. Requires a build step (already handled by Vite).')
doc.add_heading('CSS Modules / Styled Components:', level=2)
doc.add_paragraph('Pros: Scopes CSS locally to a specific component so styles never leak and affect other parts of the site.')
doc.add_paragraph('Cons: Still requires writing traditional CSS rules manually.')

# 3. Recommendation for Commercial Website
doc.add_heading('3. Recommendation for a Commercial Website', level=1)
doc.add_paragraph('For a commercial, production-ready website, I highly recommend transitioning to Tailwind CSS combined with a component library (like shadcn/ui or Material-UI).')
doc.add_paragraph('Why? Commercial applications require rapid iteration, consistent branding, and highly responsive designs. Tailwind CSS provides a standardized set of spacing, typography, and color scales out-of-the-box. It significantly speeds up development and makes it much easier for multiple developers to work on the same project without CSS conflicts.')

# 4. File-by-File Breakdown
doc.add_heading('4. Detailed File-by-File Breakdown', level=1)

files = {
    "webapp/index.html": "The main entry point of the web app. It loads the React scripts and contains the root <div> where the entire app is rendered.",
    "webapp/src/main.tsx": "The bootstrapper file. It takes the React App component and injects it into the index.html root div. It also wraps the app in the BrowserRouter for routing.",
    "webapp/src/App.tsx": "The main layout component. It contains the Sidebar, Top Header, and the Routes configuration that decides which page to show based on the URL.",
    "webapp/src/index.css": "The global design system. It contains all our CSS variables (colors, fonts, shadows) and base styles that apply to the whole app.",
    "webapp/src/pages/Dashboard.tsx": "The landing page of the app. Shows high-level statistics, emergency alerts, and recent patient activity.",
    "webapp/src/pages/AICompanion.tsx": "The chat interface where users interact with the Gemini AI for triage or fitness advice.",
    "webapp/src/pages/Fitness.tsx": "The wellness page that calculates BMI, BMR, and offers fitness routines based on user inputs.",
    "webapp/src/pages/Records.tsx": "The patient management page. Uses localForage to store records offline and allows syncing to the backend.",
    "backend/server.js": "The Node.js Express server. It handles incoming HTTP requests from the frontend, securely holds the Gemini API key, and forwards user prompts to Google's AI servers.",
    "backend/.env": "A hidden environment variable file that securely stores sensitive data, like your Google AI Studio API key. This file should never be shared publicly."
}

for filename, desc in files.items():
    p = doc.add_paragraph()
    p.add_run(f"{filename}:\n").bold = True
    p.add_run(desc)

# Save the document
save_path = os.path.join(os.path.expanduser('~'), 'Downloads', 'MedPal_Architecture_Report.docx')
doc.save(save_path)
print(f"Document saved successfully to {save_path}")
